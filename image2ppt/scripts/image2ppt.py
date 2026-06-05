import argparse
from math import cos, sin, pi

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


BLUE = "004B84"
BLUE_RGB = RGBColor.from_string(BLUE)
TEXT = "1D2733"
TEXT_RGB = RGBColor.from_string(TEXT)
LIGHT = "F3F8FC"
FONT = "微软雅黑"


def I(value):
    return Inches(value)


def set_run_font(run, size, color=TEXT, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

    # python-pptx sets latin fonts only; this explicitly sets East Asian text.
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set("typeface", FONT)


def format_text(
    shape,
    size,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.CENTER,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0.02,
):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = I(margin)
    tf.margin_right = I(margin)
    tf.margin_top = I(margin)
    tf.margin_bottom = I(margin)
    for para in tf.paragraphs:
        para.alignment = align
        para.space_after = Pt(0)
        for run in para.runs:
            set_run_font(run, size, color, bold)


def add_text(slide, text, x, y, w, h, size, color=TEXT, bold=False,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.02):
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    box.text = text
    format_text(box, size, color, bold, align, valign, margin)
    return box


def style_shape(shape, fill=None, line=BLUE, width=1.4, dash=None):
    if hasattr(shape, "shadow"):
        shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(width)
        if dash:
            shape.line.dash_style = dash
    return shape


def set_round_radius(shape, adj=6500):
    geom = shape._element.spPr.prstGeom
    avlst = geom.find(qn("a:avLst"))
    if avlst is None:
        avlst = OxmlElement("a:avLst")
        geom.append(avlst)
    for gd in list(avlst):
        if gd.get("name") == "adj":
            avlst.remove(gd)
    gd = OxmlElement("a:gd")
    gd.set("name", "adj")
    gd.set("fmla", f"val {adj}")
    avlst.append(gd)


def add_round_rect(slide, x, y, w, h, fill=LIGHT, line=BLUE, width=1.4, dash=None, radius=6500):
    shape = slide.shapes.add_shape(SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    set_round_radius(shape, radius)
    return style_shape(shape, fill, line, width, dash)


def add_rect(slide, x, y, w, h, fill=LIGHT, line=BLUE, width=1.4, dash=None):
    shape = slide.shapes.add_shape(SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    return style_shape(shape, fill, line, width, dash)


def add_line(slide, x1, y1, x2, y2, color=BLUE, width=1.5, dash=None):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    line.line.color.rgb = RGBColor.from_string(color)
    line.line.width = Pt(width)
    if dash:
        line.line.dash_style = dash
    return line


def add_triangle(slide, cx, cy, size, rotation=0, fill=BLUE):
    tri = slide.shapes.add_shape(
        SHAPE.ISOSCELES_TRIANGLE,
        I(cx - size / 2),
        I(cy - size / 2),
        I(size),
        I(size),
    )
    tri.rotation = rotation
    style_shape(tri, fill, None)
    return tri


def add_right_arrow(slide, x, y, w, h, fill=BLUE):
    arrow = slide.shapes.add_shape(SHAPE.RIGHT_ARROW, I(x), I(y), I(w), I(h))
    return style_shape(arrow, fill, None)


def add_down_arrow(slide, x, y, w, h, fill=BLUE):
    arrow = slide.shapes.add_shape(SHAPE.DOWN_ARROW, I(x), I(y), I(w), I(h))
    return style_shape(arrow, fill, None)


def add_bidirectional_dashed(slide, cx, y1, y2):
    add_line(slide, cx, y1 + 0.07, cx, y2 - 0.07, width=1.5, dash=MSO_LINE_DASH_STYLE.DASH)
    add_triangle(slide, cx, y1 + 0.04, 0.12, rotation=0)
    add_triangle(slide, cx, y2 - 0.04, 0.12, rotation=180)


def draw_cube(slide, x, y, s, color="FFFFFF", width=2.1):
    p = {
        "top": (x + 0.50 * s, y + 0.02 * s),
        "lt": (x + 0.13 * s, y + 0.22 * s),
        "lb": (x + 0.13 * s, y + 0.63 * s),
        "bottom": (x + 0.50 * s, y + 0.84 * s),
        "rb": (x + 0.87 * s, y + 0.63 * s),
        "rt": (x + 0.87 * s, y + 0.22 * s),
        "mid": (x + 0.50 * s, y + 0.42 * s),
    }
    for a, b in [
        ("top", "lt"), ("lt", "lb"), ("lb", "bottom"), ("bottom", "rb"),
        ("rb", "rt"), ("rt", "top"), ("lt", "mid"), ("mid", "rt"),
        ("mid", "bottom"),
    ]:
        add_line(slide, *p[a], *p[b], color=color, width=width)


def draw_gear(slide, cx, cy, r=0.18, color=BLUE):
    outer = slide.shapes.add_shape(SHAPE.OVAL, I(cx - r), I(cy - r), I(2 * r), I(2 * r))
    style_shape(outer, None, color, width=1.6)
    inner = slide.shapes.add_shape(SHAPE.OVAL, I(cx - r * 0.38), I(cy - r * 0.38), I(2 * r * 0.38), I(2 * r * 0.38))
    style_shape(inner, None, color, width=1.2)
    for i in range(8):
        a = i * pi / 4
        x1 = cx + cos(a) * r * 0.96
        y1 = cy + sin(a) * r * 0.96
        x2 = cx + cos(a) * r * 1.28
        y2 = cy + sin(a) * r * 1.28
        add_line(slide, x1, y1, x2, y2, color=color, width=1.7)


def draw_clipboard(slide, x, y, s=0.58, check=False):
    board = add_round_rect(slide, x + 0.07 * s, y + 0.12 * s, 0.70 * s, 0.78 * s, fill=BLUE, line=BLUE, width=1)
    clip = add_round_rect(slide, x + 0.25 * s, y + 0.04 * s, 0.35 * s, 0.18 * s, fill="FFFFFF", line="FFFFFF", width=0.6)
    for yy in (0.35, 0.50, 0.65):
        add_line(slide, x + 0.22 * s, y + yy * s, x + 0.62 * s, y + yy * s, color="FFFFFF", width=1.0)
    if check:
        add_line(slide, x + 0.25 * s, y + 0.54 * s, x + 0.37 * s, y + 0.67 * s, color="FFFFFF", width=2.0)
        add_line(slide, x + 0.37 * s, y + 0.67 * s, x + 0.62 * s, y + 0.38 * s, color="FFFFFF", width=2.0)
    return board, clip


def draw_document(slide, x, y, s=0.58):
    doc = add_round_rect(slide, x + 0.10 * s, y + 0.10 * s, 0.70 * s, 0.82 * s, fill=None, line=BLUE, width=2)
    for yy in (0.30, 0.47, 0.64):
        add_line(slide, x + 0.23 * s, y + yy * s, x + 0.65 * s, y + yy * s, width=1.7)
    return doc


def draw_person(slide, x, y, s=0.62):
    head = slide.shapes.add_shape(SHAPE.OVAL, I(x + 0.31 * s), I(y + 0.07 * s), I(0.36 * s), I(0.36 * s))
    style_shape(head, BLUE, None)
    body = slide.shapes.add_shape(SHAPE.OVAL, I(x + 0.16 * s), I(y + 0.43 * s), I(0.68 * s), I(0.50 * s))
    style_shape(body, BLUE, None)


def draw_factory(slide, x, y, s=0.78, color=BLUE):
    add_rect(slide, x + 0.02 * s, y + 0.66 * s, 0.82 * s, 0.12 * s, fill=color, line=None)
    add_rect(slide, x + 0.10 * s, y + 0.27 * s, 0.14 * s, 0.39 * s, fill=color, line=None)
    add_rect(slide, x + 0.33 * s, y + 0.08 * s, 0.14 * s, 0.58 * s, fill=color, line=None)
    add_rect(slide, x + 0.58 * s, y + 0.19 * s, 0.14 * s, 0.47 * s, fill=color, line=None)
    add_rect(slide, x + 0.00 * s, y + 0.78 * s, 0.88 * s, 0.09 * s, fill=color, line=None)
    add_line(slide, x + 0.02 * s, y + 0.66 * s, x + 0.18 * s, y + 0.53 * s, color=color, width=3.0)
    add_line(slide, x + 0.18 * s, y + 0.53 * s, x + 0.34 * s, y + 0.66 * s, color=color, width=3.0)
    add_line(slide, x + 0.34 * s, y + 0.66 * s, x + 0.50 * s, y + 0.53 * s, color=color, width=3.0)
    add_line(slide, x + 0.50 * s, y + 0.53 * s, x + 0.68 * s, y + 0.66 * s, color=color, width=3.0)


def add_extension_box(slide, x, y, w, h, suffix):
    box = add_round_rect(slide, x, y, w, h, fill="FFFFFF", line=BLUE, width=1.4, dash=MSO_LINE_DASH_STYLE.DASH)
    draw_gear(slide, x + 0.31, y + h / 2, r=0.14)
    add_text(slide, f"电厂扩展子流程\n（{suffix}）", x + 0.63, y + 0.15, w - 0.76, h - 0.30, 13.6)
    return box


def add_node(slide, x, y, w, h, label, icon_type):
    node = add_round_rect(slide, x, y, w, h, fill=LIGHT, line=BLUE, width=1.4, radius=9000)
    ix = x + 0.20
    iy = y + 0.22
    if icon_type == "clipboard":
        draw_clipboard(slide, ix, iy, s=0.62)
    elif icon_type == "document":
        draw_document(slide, ix, iy, s=0.64)
    elif icon_type == "person":
        draw_person(slide, ix, iy, s=0.68)
    elif icon_type == "check":
        draw_clipboard(slide, ix, iy, s=0.62, check=True)
    add_text(slide, label, x + 0.78, y + 0.24, w - 0.90, h - 0.46, 19.5, color=BLUE, bold=True)
    return node


def add_plant_card(slide, y, name):
    card = add_round_rect(slide, 11.35, y, 4.20, 2.40, fill="F8FBFE", line=BLUE, width=1.0, radius=4500)
    draw_factory(slide, 11.55, y + 0.28, s=0.70)
    add_text(slide, name, 12.25, y + 0.36, 1.35, 0.42, 18.0, color=BLUE, bold=True, align=PP_ALIGN.LEFT)

    left = add_round_rect(slide, 11.55, y + 1.05, 1.15, 0.58, fill="FFFFFF", line=BLUE, width=1.3, dash=MSO_LINE_DASH_STYLE.DASH)
    add_text(slide, "前置扩展", 11.59, y + 1.18, 1.07, 0.28, 11.8, color=BLUE, bold=True)
    add_right_arrow(slide, 12.72, y + 1.23, 0.34, 0.22)

    mid = add_round_rect(slide, 13.08, y + 1.05, 0.98, 0.58, fill=BLUE, line=BLUE, width=1.0)
    add_text(slide, "节点X", 13.15, y + 1.17, 0.84, 0.30, 12.6, color="FFFFFF", bold=True)
    add_right_arrow(slide, 14.08, y + 1.23, 0.34, 0.22)

    right = add_round_rect(slide, 14.44, y + 1.05, 1.05, 0.58, fill="FFFFFF", line=BLUE, width=1.3, dash=MSO_LINE_DASH_STYLE.DASH)
    add_text(slide, "后置扩展", 14.47, y + 1.18, 0.99, 0.28, 11.8, color=BLUE, bold=True)
    add_text(slide, "可配置任意扩展子流程", 12.28, y + 1.87, 2.45, 0.33, 12.8, color=TEXT)
    return card


def build_deck():
    prs = Presentation()
    prs.slide_width = I(15.84)
    prs.slide_height = I(10.24)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")

    # Standard product entry.
    product = add_round_rect(slide, 0.51, 1.28, 3.00, 1.31, fill=BLUE, line=BLUE, width=1, radius=6500)
    draw_cube(slide, 0.84, 1.57, 0.73, color="FFFFFF", width=2.4)
    add_text(slide, "标准产品", 1.68, 1.72, 1.64, 0.48, 21.5, color="FFFFFF", bold=True)
    add_down_arrow(slide, 1.77, 2.63, 0.50, 0.55)

    note = add_round_rect(slide, 4.28, 1.22, 3.10, 1.70, fill="FFFFFF", line=BLUE, width=1.3, dash=MSO_LINE_DASH_STYLE.DASH, radius=5000)
    add_text(
        slide,
        "标准产品提供标准流程，\n各电厂可在每个节点前后\n插入扩展子流程，满足个\n性化审批扩展需求。",
        4.50,
        1.42,
        2.70,
        1.27,
        14.5,
        align=PP_ALIGN.LEFT,
        margin=0.01,
    )

    # Main standard flow area.
    main = add_round_rect(slide, 0.31, 3.16, 10.40, 6.20, fill="FFFFFF", line=BLUE, width=1.3, radius=2500)
    header = add_rect(slide, 0.31, 3.16, 10.40, 0.58, fill=BLUE, line=BLUE, width=1)
    add_text(slide, "标准流程", 4.43, 3.23, 2.15, 0.40, 21.0, color="FFFFFF", bold=True)

    top_y, bot_y = 4.07, 7.70
    ext_w, ext_h = 2.28, 1.00
    node_y, node_h = 5.90, 1.03
    specs = [
        (0.48, 0.56, 1.98, "节点1", "clipboard"),
        (3.20, 3.25, 2.12, "节点2", "document"),
        (5.82, 5.84, 2.15, "节点3", "person"),
        (8.39, 8.45, 1.94, "节点4", "check"),
    ]
    node_centers = []
    for ext_x, node_x, node_w, label, icon in specs:
        add_extension_box(slide, ext_x, top_y, ext_w, ext_h, "前置")
        add_extension_box(slide, ext_x, bot_y, ext_w, ext_h, "后置")
        add_node(slide, node_x, node_y, node_w, node_h, label, icon)
        cx = node_x + node_w / 2
        node_centers.append(cx)
        add_bidirectional_dashed(slide, cx, top_y + ext_h, node_y)
        add_bidirectional_dashed(slide, cx, node_y + node_h, bot_y)

    add_right_arrow(slide, 2.54, 6.20, 0.62, 0.42)
    add_right_arrow(slide, 5.38, 6.20, 0.45, 0.42)
    add_right_arrow(slide, 8.04, 6.20, 0.35, 0.42)
    add_right_arrow(slide, 10.38, 6.15, 0.83, 0.52)

    # Right customization panel.
    panel = add_round_rect(slide, 11.18, 0.25, 4.50, 9.98, fill="FFFFFF", line=BLUE, width=1.5, radius=2500)
    panel_header = add_round_rect(slide, 11.18, 0.25, 4.50, 0.75, fill=BLUE, line=BLUE, width=1, radius=4500)
    add_rect(slide, 11.18, 0.75, 4.50, 0.25, fill=BLUE, line=BLUE, width=0)
    add_text(slide, "各电厂个性化审批扩展", 11.82, 0.43, 3.25, 0.35, 20.0, color="FFFFFF", bold=True)
    add_text(
        slide,
        "各电厂可在任意节点的前置或后\n置位置，配置个性化审批扩展。",
        11.55,
        1.18,
        3.72,
        0.85,
        15.3,
        color=TEXT,
        align=PP_ALIGN.LEFT,
        margin=0,
    )

    add_plant_card(slide, 2.24, "电厂A")
    add_plant_card(slide, 4.84, "电厂B")
    add_plant_card(slide, 7.40, "电厂C")
    add_text(slide, "……", 13.22, 9.92, 0.55, 0.20, 14.0, color=TEXT)

    return prs


def main():
    parser = argparse.ArgumentParser(description="Generate an editable PPTX deck from this image2ppt scaffold.")
    parser.add_argument("-o", "--output", default="approval-extension-flow-editable.pptx", help="Output PPTX path.")
    args = parser.parse_args()

    deck = build_deck()
    deck.save(args.output)


if __name__ == "__main__":
    main()
